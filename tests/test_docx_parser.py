import sys
import os
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
from app.core.doc_parser import extract_text_from_file, extract_relevant_marking_scheme

def test_docx_rubric_extraction():
    print("Testing Word Document (.docx) Rubric Parsing...")
    
    # 1. Create a mock Word (.docx) document with boilerplate and rubrics
    test_docx_path = Path("tests/mock_marking_scheme.docx")
    doc = docx.Document()
    
    # Administrative boilerplate
    doc.add_heading("MID-YEAR EXAMINATION 2026", level=1)
    doc.add_paragraph("INSTRUCTIONS TO CANDIDATES:")
    doc.add_paragraph("1. Do not open this examination booklet until you are told to do so.")
    doc.add_paragraph("2. Write your full name and student index number on the cover sheet.")
    doc.add_paragraph("3. Answer all questions in black or dark blue ballpoint pen.")
    doc.add_paragraph("4. Total time allowed: 2 hours.")
    doc.add_paragraph("--------------------------------------------------")
    
    # Actual Marking Scheme
    doc.add_heading("OFFICIAL MARKING SCHEME & RUBRIC", level=2)
    doc.add_paragraph("Subject: Physics - Year 11 Mechanics")
    doc.add_paragraph("Total Marks: 50")
    
    # Table of Question criteria
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Question"
    hdr_cells[1].text = "Criteria / Model Answer"
    hdr_cells[2].text = "Marks"
    
    # Add rows
    row_1 = table.add_row().cells
    row_1[0].text = "Q1"
    row_1[1].text = "State Newton's Second Law: F = ma. Award 2 marks for formula, 3 marks for correct explanation."
    row_1[2].text = "5"
    
    row_2 = table.add_row().cells
    row_2[0].text = "Q2"
    row_2[1].text = "Calculate kinetic energy: Ek = 0.5 * m * v^2. Award 3 marks for substitution, 2 marks for answer with Joules."
    row_2[2].text = "5"
    
    doc.save(str(test_docx_path))
    
    # 2. Extract text and table from .docx
    print("Extracting text and tables from .docx...", flush=True)
    extracted_text = extract_text_from_file(str(test_docx_path))
    assert "MID-YEAR EXAMINATION" in extracted_text
    assert "Newton's Second Law" in extracted_text
    assert "kinetic energy" in extracted_text
    print("DOCX text & table extraction passed!", flush=True)
    
    # 3. Test relevant marking scheme isolation
    print("Isolating relevant marking scheme portion...", flush=True)
    isolated = extract_relevant_marking_scheme(extracted_text)
    assert len(isolated["clean_marking_scheme"]) > 0
    print("Isolated Marking Scheme Preview:\n", isolated["clean_marking_scheme"][:200], flush=True)
    
    # Clean up test file
    if test_docx_path.exists():
        test_docx_path.unlink()
        
    print("Word docx rubric testing passed successfully!")

if __name__ == "__main__":
    test_docx_rubric_extraction()
